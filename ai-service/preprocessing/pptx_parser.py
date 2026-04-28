"""
PPTX 文本提取器，逐页逐 shape 提取文本。
"""

from .parser import BaseParser


class PptxParser(BaseParser):
    """使用 python-pptx 提取 .pptx 文件中的文本。"""

    def extract_text(self, file_path: str) -> str:
        from pptx import Presentation

        prs = Presentation(file_path)
        parts: list[str] = []
        for slide in prs.slides:
            slide_texts: list[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        text = para.text.strip()
                        if text:
                            slide_texts.append(text)
            if slide_texts:
                parts.append("\n".join(slide_texts))
        return "\n\n---\n\n".join(parts)
